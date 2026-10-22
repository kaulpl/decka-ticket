from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1] / "outputs" / "Regulamin-sprzedazy-biletow-Decka-Pelplin-propozycja.docx"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def add_clause(doc, number, title, paragraphs, bullets=None):
    doc.add_heading(f"§ {number}. {title}", level=1)
    for text in paragraphs:
        p = doc.add_paragraph(text)
        p.paragraph_format.space_after = Pt(6)
    for text in bullets or []:
        p = doc.add_paragraph(text, style="List Bullet")
        p.paragraph_format.space_after = Pt(4)

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.line_spacing = 1.12
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(26)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(20, 43, 60)
for name, size in [("Heading 1", 15), ("Heading 2", 12)]:
    styles[name].font.name = "Aptos Display"
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(20, 43, 60)

header = section.header.paragraphs[0]
header.text = "DECKA PELPLIN  |  PROJEKT REGULAMINU"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor(90, 105, 118)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Projekt roboczy • stan prawny sprawdzony 22 października 2026 r. • wymaga uzupełnienia danych Organizatora")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(90, 105, 118)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("REGULAMIN SPRZEDAŻY BILETÓW\nNA WYDARZENIA SPORTOWE")
sub = doc.add_paragraph("Decka Pelplin — propozycja do wdrożenia", style="Subtitle")
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER

table = doc.add_table(rows=4, cols=2)
table.style = "Table Grid"
labels = ["Organizator / Sprzedawca", "Adres i dane rejestrowe", "Kontakt reklamacyjny", "Sklep internetowy"]
values = ["[pełna nazwa podmiotu prowadzącego klub]", "[adres, NIP, KRS/ewidencja]", "[e-mail i adres korespondencyjny]", "[adres strony sprzedaży biletów]"]
for i, (label, value) in enumerate(zip(labels, values)):
    table.cell(i, 0).text = label
    table.cell(i, 1).text = value
    shade(table.cell(i, 0), "EAF0F4")
    table.cell(i, 0).paragraphs[0].runs[0].bold = True

note = doc.add_table(rows=1, cols=1)
note.style = "Table Grid"
note.cell(0, 0).text = ("WAŻNE PRZED PUBLIKACJĄ: uzupełnij pola w nawiasach kwadratowych, wskaż poprawny podmiot zawierający umowę z klientem oraz e-mail reklamacyjny. "
                         "Dokument jest praktycznym projektem opartym na aktualnych przepisach, ale przy publikacji regulaminu konsumenckiego zalecana jest końcowa weryfikacja przez prawnika znącego model sprzedaży klubu.")
shade(note.cell(0, 0), "FFF4CC")

add_clause(doc, 1, "Postanowienia ogólne", [
    "1. Niniejszy regulamin określa zasady sprzedaży, wydawania i korzystania z biletów na mecze oraz inne wydarzenia sportowe organizowane lub sprzedawane przez Organizatora.",
    "2. Regulamin jest udostępniany nieodpłatnie przed złożeniem zamówienia w sposób umożliwiający jego zapisanie i odtworzenie. Klient powinien zapoznać się z nim przed zakupem.",
    "3. W zakresie nieuregulowanym Regulaminem stosuje się przepisy prawa polskiego, w szczególności Kodeksu cywilnego oraz ustawy o prawach konsumenta. Postanowienia Regulaminu nie wyłączają ani nie ograniczają bezwzględnie obowiązujących praw konsumenta."
])

add_clause(doc, 2, "Definicje", [], [
    "Bilet — dokument elektroniczny, w szczególności plik PDF lub zapis w portfelu mobilnym, zawierający unikalny kod i uprawniający do jednorazowego wstępu na wskazane Wydarzenie oraz zajęcia oznaczonego miejsca.",
    "Klient — osoba składająca zamówienie; Konsument — klient będący konsumentem w rozumieniu Kodeksu cywilnego albo osobą korzystającą z odpowiedniej ochrony konsumenckiej na podstawie ustawy.",
    "Wydarzenie — mecz lub inne wydarzenie sportowe oznaczone datą, miejscem i zakresem uprawnień wynikających z Biletu.",
    "Siła Wyższa — zewnętrzne, nadzwyczajne i niezależne od Organizatora zdarzenie, którego skutkom nie można było rozsądnie zapobiec, w tym decyzje władz lub organizatora rozgrywek, zagrożenia bezpieczeństwa, awarie obiektu lub infrastruktury, klęski żywiołowe, epidemie, działania wojenne albo poważne zakłócenia porządku publicznego."
])

add_clause(doc, 3, "Składanie zamówienia i płatność", [
    "1. Informacje prezentowane w systemie sprzedaży stanowią zaproszenie do zawarcia umowy. Klient wybiera Wydarzenie, miejsce, rodzaj Biletu i podaje dane niezbędne do realizacji zamówienia.",
    "2. Klient odpowiada za poprawność podanego adresu e-mail i innych danych. Przed opłaceniem zamówienia system wskazuje łączną cenę oraz istotne cechy świadczenia.",
    "3. Umowa zostaje zawarta po skutecznym potwierdzeniu płatności przez operatora płatności albo po wydaniu Biletu przez uprawnionego kasjera. Sam powrót klienta ze strony operatora płatności nie stanowi potwierdzenia zapłaty.",
    "4. Nieopłacone zamówienie może wygasnąć po czasie wskazanym w systemie, a zablokowane miejsca wracają wtedy do sprzedaży."
])

add_clause(doc, 4, "Bilet i zasady wstępu", [
    "1. Każdy Bilet zawiera unikalny kod. Pierwsze prawidłowe zeskanowanie kodu powoduje jego oznaczenie jako wykorzystanego. Udostępnienie kodu osobie trzeciej odbywa się na ryzyko posiadacza Biletu.",
    "2. Organizator może odmówić wstępu na podstawie kodu nieważnego, wykorzystanego, unieważnionego, podrobionego albo dotyczącego innego Wydarzenia.",
    "3. Posiadacz Biletu przestrzega regulaminu obiektu i zasad bezpieczeństwa. Bilet ulgowy wymaga okazania dokumentu potwierdzającego uprawnienie do ulgi, jeżeli Organizator tego zażąda. Brak dokumentu może skutkować koniecznością dopłaty do Biletu normalnego, jeżeli taka możliwość jest dostępna."
])

add_clause(doc, 5, "Brak dobrowolnych zwrotów i prawo odstąpienia", [
    "1. Z zastrzeżeniem § 7, prawidłowo zakupione Bilety nie podlegają dobrowolnemu zwrotowi ani wymianie na pieniądze z przyczyn leżących po stronie Klienta, w szczególności z powodu zmiany planów, choroby, spóźnienia, niewykorzystania Biletu albo błędnego wyboru Wydarzenia.",
    "2. W przypadku umowy dotyczącej wydarzenia sportowego odbywającego się w oznaczonym dniu lub okresie Konsumentowi nie przysługuje ustawowe czternastodniowe prawo odstąpienia od umowy zawartej na odległość — zgodnie z art. 38 pkt 12 ustawy o prawach konsumenta.",
    "3. Powyższe nie narusza praw Klienta w razie odwołania Wydarzenia, jego istotnej zmiany, niewykonania umowy lub niezgodności świadczenia z umową."
])

add_clause(doc, 6, "Zmiana miejsca", [
    "1. Na wniosek Klienta Organizator może, w miarę dostępności i możliwości technicznych, dokonać jednorazowej zmiany miejsca na to samo Wydarzenie. Zmiana miejsca nie stanowi zwrotu Biletu.",
    "2. Organizator może uzależnić zmianę od dopłaty różnicy ceny. Różnica na korzyść Klienta nie podlega wypłacie, jeżeli Klient dobrowolnie zaakceptuje miejsce tańsze po otrzymaniu jasnej informacji o skutku zmiany.",
    "3. Po zmianie miejsca poprzedni kod Biletu zostaje unieważniony, a wiążący jest nowo wydany Bilet. Organizator może odmówić zmiany, jeżeli brak jest odpowiednich wolnych miejsc albo wniosek złożono zbyt późno, aby bezpiecznie go obsłużyć."
])

add_clause(doc, 7, "Zmiana terminu, miejsca lub odwołanie Wydarzenia", [
    "1. Terminy rozgrywek mogą ulec zmianie, w szczególności na podstawie decyzji ligi, związku sportowego, właściciela obiektu, organu władzy lub ze względów bezpieczeństwa. Organizator poinformuje o znanej zmianie za pomocą strony internetowej, systemu biletowego lub danych kontaktowych Klienta.",
    "2. W razie przełożenia Wydarzenia Bilet zachowuje ważność na nowy termin. Jeżeli zmiana jest istotna, Konsument może nie zaakceptować nowego terminu i skorzystać z praw przysługujących mu na podstawie obowiązujących przepisów, w tym żądać zwrotu ceny za niewykonane świadczenie.",
    "3. W razie całkowitego odwołania Wydarzenia Klientowi przysługuje zwrot zapłaconej ceny Biletu oraz innych obowiązkowych opłat pobranych przy jego zakupie, w zakresie wymaganym przez prawo. Organizator może zaproponować voucher, zachowanie Biletu na nowy termin lub inne świadczenie zastępcze, lecz przyjęcie takiej propozycji jest dobrowolne.",
    "4. Jeżeli odwołanie, przełożenie lub zmiana Wydarzenia wynika z Siły Wyższej albo okoliczności, za które Organizator nie ponosi odpowiedzialności, Organizator — w najszerszym zakresie dozwolonym prawem — nie odpowiada za pośrednie koszty Klienta, takie jak podróż, nocleg, wyżywienie, utracony zarobek czy inne wydatki niezwiązane bezpośrednio z ceną Biletu. Postanowienie to nie wyłącza obowiązku zwrotu ceny za niewykonane świadczenie ani odpowiedzialności, której nie można wyłączyć na mocy prawa.",
    "5. Organizator nie wyłącza odpowiedzialności za szkodę wyrządzoną umyślnie ani innych przypadków, w których ograniczenie odpowiedzialności byłoby nieważne."
])

add_clause(doc, 8, "Reklamacje", [
    "1. Reklamację dotyczącą sprzedaży lub działania Biletu można złożyć na adres [e-mail reklamacyjny] albo pisemnie na adres Organizatora.",
    "2. Wskazane jest podanie numeru zamówienia, danych kontaktowych, opisu problemu i oczekiwanego sposobu rozwiązania. Brak tych danych nie pozbawia Konsumenta praw, lecz może utrudnić rozpoznanie sprawy.",
    "3. Organizator udzieli odpowiedzi Konsumentowi w terminie 14 dni od otrzymania reklamacji, chyba że bezwzględnie obowiązujący przepis przewiduje inny termin."
])

add_clause(doc, 9, "Dane osobowe i komunikacja", [
    "1. Dane osobowe są przetwarzane w celu zawarcia i wykonania umowy, obsługi płatności, dostarczenia Biletu, obsługi wejścia, reklamacji oraz wykonania obowiązków prawnych. Szczegółowe informacje zawiera polityka prywatności Organizatora.",
    "2. Komunikaty niezbędne do realizacji umowy, w tym o zmianie terminu lub statusie zamówienia, nie stanowią informacji handlowej wymagającej odrębnej zgody marketingowej."
])

add_clause(doc, 10, "Postanowienia końcowe", [
    "1. Regulamin obowiązuje od [data]. Do danego zamówienia stosuje się wersję udostępnioną Klientowi przed jego złożeniem.",
    "2. Organizator może zmienić Regulamin z ważnych przyczyn prawnych, organizacyjnych lub technicznych. Zmiana nie narusza praw nabytych i nie pogarsza warunków już zawartej umowy, chyba że wynika to z bezwzględnie obowiązującego prawa.",
    "3. Ewentualna nieważność pojedynczego postanowienia nie narusza pozostałych postanowień. W relacjach z Konsumentami w miejsce nieważnego postanowienia stosuje się przepisy prawa."
])

doc.add_page_break()
doc.add_heading("Notatka wdrożeniowa i podstawa prawna", level=1)
doc.add_paragraph("Ta strona nie musi być publikowana jako część regulaminu. Służy osobie wdrażającej dokument.")
for text in [
    "Ustawa o prawach konsumenta, art. 38 pkt 12 — wyłączenie 14-dniowego prawa odstąpienia dla usług związanych z wydarzeniami sportowymi, jeżeli oznaczono dzień lub okres świadczenia.",
    "Kodeks cywilny, art. 471–475 i art. 494–495 — odpowiedzialność za niewykonanie zobowiązania, granice wyłączeń oraz obowiązek zwrotu otrzymanego świadczenia, gdy wzajemne świadczenie stało się niemożliwe z przyczyn niezawinionych.",
    "Stanowisko UOKiK dotyczące odwołanych wydarzeń — konsument może odrzucić voucher lub inne świadczenie zastępcze i żądać zwrotu ceny; nie należy uzależniać ustawowego zwrotu od krótkiego, pozaustawowego terminu.",
    "Przed publikacją należy spójnie uzupełnić: dane sprzedawcy, kanał reklamacyjny, datę wejścia w życie, politykę prywatności i regulamin obiektu."
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("Źródła urzędowe", level=2)
sources = [
    "Ustawa z 30 maja 2014 r. o prawach konsumenta (tekst w ELI): https://eli.gov.pl/api/acts/DU/2020/287/text.html",
    "Kodeks cywilny (ELI; art. 471–475, 494–495): https://eli.gov.pl/api/acts/DU/2023/1610/text.html",
    "UOKiK, Poradnik koncertowy 2026 — odwołane i przełożone wydarzenia: https://uokik.gov.pl/Download/1882",
    "UOKiK, decyzja RGD-2/2026 — niedozwolone ograniczanie zwrotu ceny za odwołane wydarzenie: https://uokik.gov.pl/Download/1880",
]
for source in sources:
    doc.add_paragraph(source, style="List Bullet")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
